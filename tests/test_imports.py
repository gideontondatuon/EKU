"""Actual supplied workbook plus explicitly synthetic edge-case workbooks."""

import io
import zipfile
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pytest
from django.contrib.auth import get_user_model
from django.core.exceptions import PermissionDenied, ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import close_old_connections
from openpyxl import Workbook

from cashflow.models import CashflowTransaction
from cashflow.services import archive_transactions, model_fingerprint
from core.models import AuditLog, Denomination, SystemSettings
from imports.models import ImportBatch, ImportRow, VerificationDecision
from imports.parser import ADAPTERS, DUMMY_CHECKSUM, WorkbookError, parse_workbook
from imports.services import commit_batch, commit_replace, decide_row, preview_replace, stage_import

DUMMY = Path(__file__).resolve().parents[1] / "references" / "SETBAY_DUMMY.xlsx"
SCOPE = {"start": "2021-09-01", "end": "2021-09-30", "region_codes": ["MANADO"]}


@pytest.fixture(autouse=True)
def private_uploads(settings, tmp_path):
    settings.MEDIA_ROOT = tmp_path / "private"


def source_rows(*overrides, **base_changes):
    base = {
        "transaction_date": datetime(2021, 9, 10),
        "transaction_type": "Deposit",
        "money_type": "D-SB ULE",
        "bank_name": "BANK SINTETIS A",
        "account": "0000000001",
        "branch_code": 760,
        "branch_name": "KPwBI Provinsi Sulawesi Utara",
        "denomination_code": "SYN100",
        "denomination_value": "100.000",
        "denomination_order": 1,
        "source_amount": 1000000,
        "pieces": 10,
        "inflow_amount": 1000000,
        "outflow_amount": 0,
        "source_netflow": 1000000,
    }
    return [{**base, **base_changes, **changes} for changes in (overrides or ({},))]


def workbook_bytes(rows=None, marker="synthetic", date1904=False):
    book = Workbook()
    sheet = book.active
    sheet.title = "SETBAY MANADO"
    if date1904:
        from openpyxl.utils.datetime import CALENDAR_MAC_1904

        book.epoch = CALENDAR_MAC_1904
    keys = list(ADAPTERS[sheet.title])
    sheet.append([ADAPTERS[sheet.title][key][0] for key in keys])
    for row in rows or source_rows():
        sheet.append([row[key] for key in keys])
    refs = book.create_sheet("DENOM")
    refs.append(["Pecahan", "Denom", "Urutan Denom", None, "Kode Kastip", "Nama Kastip"])
    refs.append(["SYN100", "100.000", 1, None, "000001", "SKM SINTETIS"])
    book.create_sheet("Metadata").append([marker, "DATA SINTETIS UNTUK TES"])
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def upload(content, name="synthetic.xlsx"):
    return SimpleUploadedFile(
        name,
        content,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def rewrite_zip(data, replacements=None, extras=None):
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as original, zipfile.ZipFile(output, "w") as target:
        for entry in original.infolist():
            value = original.read(entry)
            if replacements and entry.filename in replacements:
                value = replacements[entry.filename](value)
            target.writestr(entry, value)
        for name, value in (extras or {}).items():
            target.writestr(name, value)
    return output.getvalue()


def test_actual_workbook_cells_dates_cached_formula_and_quality():
    parsed = parse_workbook(DUMMY.read_bytes(), DUMMY.name)
    assert parsed.metadata["checksum"] == DUMMY_CHECKSUM
    assert parsed.metadata["source_update_date"] == "2026-03-26"
    assert Counter(row.status for row in parsed.rows) == {"NEW_VALID": 6, "NEEDS_REVIEW": 2}
    assert {row.normalized["transaction_date"] for row in parsed.rows} == {
        "2021-09-10",
        "2021-09-15",
    }
    assert all(row.normalized["account"].startswith("0") for row in parsed.rows)
    assert parsed.region_references["000001"] == "SKM Siau"
    following = next(r for r in parsed.rows if r.sheet == "SETBAY MANADO" and r.row_number == 3)
    assert following.raw["inflow_amount"]["formula"]["attributes"]["t"] == "shared"
    assert following.normalized["inflow_amount"] == "1000000000"
    assert sum(Decimal(r.normalized["inflow_amount"]) for r in parsed.rows) == Decimal("8254000000")


@pytest.mark.parametrize("negative", [False, True])
def test_synthetic_withdrawal_sign_conventions(negative):
    raw = workbook_bytes(
        source_rows(
            transaction_type="Withdrawal",
            inflow_amount=0,
            outflow_amount=-1000000 if negative else 1000000,
            source_netflow=-1000000,
        )
    )
    row = parse_workbook(raw, "synthetic.xlsx").rows[0]
    assert row.status == "NEW_VALID"
    assert row.normalized["outflow_amount"] == "1000000"
    assert bool(row.normalized["normalization_notes"]) == negative


def test_parser_1904_epoch_and_exact_large_decimal():
    raw = workbook_bytes(
        source_rows(
            source_amount="9007199254800000",
            pieces="90071992548",
            inflow_amount="9007199254800000",
            source_netflow="9007199254800000",
        ),
        date1904=True,
    )
    row = parse_workbook(raw, "synthetic.xlsx").rows[0]
    assert row.status == "NEW_VALID"
    assert row.normalized["inflow_amount"] == "9007199254800000"
    assert row.normalized["transaction_date"] == "2021-09-10"


@pytest.mark.parametrize(
    "change,expected",
    [
        ({"account": 12345}, "NEW_VALID"),
        ({"pieces": 9}, "NEEDS_REVIEW"),
        ({"source_netflow": 1}, "NEEDS_REVIEW"),
        ({"source_amount": "NaN"}, "INVALID"),
        ({"transaction_date": "not-a-date"}, "INVALID"),
        ({"inflow_amount": "=1+1"}, "INVALID"),
        ({"source_amount": "1.5"}, "INVALID"),
    ],
)
def test_validation_never_silently_coerces_bad_data(change, expected):
    assert (
        parse_workbook(workbook_bytes(source_rows(change)), "synthetic.xlsx").rows[0].status
        == expected
    )


def test_upload_guards_and_xml_entity_rejection():
    with pytest.raises(WorkbookError):
        parse_workbook(b"not zip", "file.xlsx")
    with pytest.raises(WorkbookError):
        parse_workbook(workbook_bytes(), "file.xlsm")
    with pytest.raises(WorkbookError):
        parse_workbook(b"x" * (10 * 1024 * 1024 + 1), "large.xlsx")
    with pytest.raises(WorkbookError):
        parse_workbook(DUMMY.read_bytes(), DUMMY.name, production=True)
    for extras in ({"xl/vbaProject.bin": b"macro"}, {"../unsafe": b"bad"}):
        with pytest.raises(WorkbookError):
            parse_workbook(rewrite_zip(workbook_bytes(), extras=extras), "synthetic.xlsx")
    xml = b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY evil SYSTEM "file:///etc/passwd">]><x>&evil;</x>'
    with pytest.raises(WorkbookError):
        parse_workbook(
            rewrite_zip(workbook_bytes(), {"xl/workbook.xml": lambda value: xml}), "synthetic.xlsx"
        )
    with pytest.raises(WorkbookError):
        parse_workbook(workbook_bytes(source_rows({}, {})), "synthetic.xlsx", row_limit=1)


@pytest.mark.django_db
def test_dummy_partial_import_and_checksum_idempotence(superuser):
    batch = stage_import(superuser, upload(DUMMY.read_bytes(), DUMMY.name))
    assert batch.is_demo and batch.stats == {"NEW_VALID": 6, "NEEDS_REVIEW": 2}
    commit_batch(superuser, batch.pk)
    assert CashflowTransaction.objects.count() == 6
    batch.refresh_from_db()
    assert batch.status == "PARTIAL"
    again = stage_import(superuser, upload(DUMMY.read_bytes(), "renamed.xlsx"))
    assert again.pk == batch.pk and ImportBatch.objects.count() == 1
    commit_batch(superuser, again.pk)
    assert CashflowTransaction.objects.count() == 6
    assert not SystemSettings.load().current_publication_version_id
    tx = CashflowTransaction.objects.first()
    assert tx.fingerprint == model_fingerprint(tx)
    assert tx.account not in str(list(AuditLog.objects.values("before", "after", "reason")))
    assert "/.." not in batch.source_file.name and batch.source_file.name.endswith(".xlsx")


@pytest.mark.django_db
def test_operational_new_master_resolves_from_consistent_reference(superuser):
    batch = stage_import(superuser, upload(workbook_bytes()))
    assert batch.stats == {"NEW_VALID": 1}
    commit_batch(superuser, batch.pk)
    assert CashflowTransaction.objects.count() == 1
    assert Denomination.objects.filter(code="SYN100", verified=True).exists()
    assert AuditLog.objects.filter(action="IMPORT_DENOMINATION_CREATE").exists()


@pytest.mark.django_db
def test_demo_master_does_not_authorize_operational_import(superuser):
    stage_import(superuser, upload(workbook_bytes()), is_demo=True)
    operational = stage_import(superuser, upload(workbook_bytes(marker="operational")))
    assert operational.stats == {"NEEDS_REVIEW": 1}
    with pytest.raises(ValidationError):
        decide_row(superuser, operational.rows.get().pk, "CREATE_NEW", "Masih master demo")


@pytest.mark.django_db
def test_invalid_and_ambiguous_do_not_block_valid_partial_rows(superuser):
    batch = stage_import(
        superuser,
        upload(
            workbook_bytes(
                source_rows(
                    {},
                    {"bank_name": "BANK B", "pieces": 9},
                    {"bank_name": "BANK C", "transaction_date": "invalid"},
                )
            )
        ),
        is_demo=True,
    )
    assert batch.stats == {"NEW_VALID": 1, "NEEDS_REVIEW": 1, "INVALID": 1}
    commit_batch(superuser, batch.pk)
    assert CashflowTransaction.objects.count() == 1


@pytest.mark.django_db
def test_review_permission_and_unresolved_conflict_cannot_be_bypassed(superuser, staff, grant):
    with pytest.raises(PermissionDenied):
        stage_import(staff, upload(DUMMY.read_bytes()))
    batch = stage_import(superuser, upload(DUMMY.read_bytes()))
    row = batch.rows.get(sheet_name="SETBAY KASTIP", source_row_number=4)
    with pytest.raises(PermissionDenied):
        decide_row(staff, row.pk, "CREATE_NEW", "Konfirmasi")
    grant(staff, "verify_ambiguous")
    with pytest.raises(ValidationError):
        decide_row(staff, row.pk, "CREATE_NEW", "Tidak boleh bypass konflik")
    Denomination.objects.create(
        code="KX16", value=50000, verified=True, source_note="Fixture sintetis approval"
    )
    decided = decide_row(
        staff,
        row.pk,
        "CREATE_NEW",
        "Koreksi dengan referensi yang disetujui",
        corrections={"denomination_value": "50000"},
    )
    assert decided.status == "IMPORTED"
    assert decided.linked_transaction.inflow_amount == 50000000
    assert VerificationDecision.objects.count() == 1


@pytest.mark.django_db
def test_real_identical_occurrences_preserved_after_explicit_review(superuser):
    batch = stage_import(superuser, upload(workbook_bytes(source_rows({}, {}))), is_demo=True)
    assert batch.stats == {"NEW_VALID": 1, "NEEDS_REVIEW": 1}
    commit_batch(superuser, batch.pk)
    pending = batch.rows.get(status="NEEDS_REVIEW")
    decide_row(superuser, pending.pk, "CREATE_NEW", "Dua kejadian berbeda dikonfirmasi operator")
    assert CashflowTransaction.objects.count() == 2
    assert len(set(CashflowTransaction.objects.values_list("fingerprint", flat=True))) == 1


@pytest.mark.django_db
def test_historical_snapshot_occurrence_reconciliation(superuser):
    first = stage_import(
        superuser,
        upload(workbook_bytes(source_rows({}, {}))),
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    commit_batch(superuser, first.pk)
    decide_row(
        superuser,
        first.rows.get(status="NEEDS_REVIEW").pk,
        "CREATE_NEW",
        "Kejadian kedua dikonfirmasi",
    )
    second = stage_import(
        superuser,
        upload(workbook_bytes(source_rows({}, {}, {"bank_name": "BANK BARU"}), marker="next")),
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    assert second.stats == {"ALREADY_IMPORTED": 2, "NEW_VALID": 1}
    commit_batch(superuser, second.pk)
    assert CashflowTransaction.objects.count() == 3
    excess = stage_import(
        superuser,
        upload(workbook_bytes(source_rows({}, {}, {}), marker="excess")),
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    assert excess.stats == {"ALREADY_IMPORTED": 2, "NEEDS_REVIEW": 1}


@pytest.mark.django_db
def test_unknown_historical_scope_requires_review_and_deleted_stays_deleted(superuser):
    first = stage_import(superuser, upload(workbook_bytes()), is_demo=True)
    commit_batch(superuser, first.pk)
    tx = CashflowTransaction.objects.get()
    archive_transactions(superuser, [tx.pk], "Arsip untuk tes")
    again = stage_import(superuser, upload(workbook_bytes(marker="repeat")), is_demo=True)
    assert again.stats == {"NEEDS_REVIEW": 1}
    row = again.rows.get()
    with pytest.raises(ValidationError):
        decide_row(superuser, row.pk, "SKIP_AS_EXISTING", "UUID wajib")
    decide_row(superuser, row.pk, "SKIP_AS_EXISTING", "Tetap arsip", target_id=tx.pk)
    assert CashflowTransaction.objects.count() == 1
    tx.refresh_from_db()
    assert tx.soft_deleted_at is not None


@pytest.mark.django_db
def test_update_requires_target_and_edit_permission(superuser, staff, grant):
    first = stage_import(superuser, upload(workbook_bytes()), is_demo=True)
    commit_batch(superuser, first.pk)
    original = CashflowTransaction.objects.get()
    modified = source_rows(
        source_amount=2000000, pieces=20, inflow_amount=2000000, source_netflow=2000000
    )
    batch = stage_import(superuser, upload(workbook_bytes(modified)), mode="UPDATE", is_demo=True)
    row = batch.rows.get()
    assert row.status == "NEEDS_REVIEW" and row.matches.count() == 1
    grant(staff, "verify_ambiguous")
    with pytest.raises(ValidationError):
        decide_row(staff, row.pk, "UPDATE_SELECTED_EXISTING", "Target wajib")
    with pytest.raises(PermissionDenied):
        decide_row(
            staff, row.pk, "UPDATE_SELECTED_EXISTING", "Izin edit wajib", target_id=original.pk
        )
    grant(staff, "edit_transaction")
    decide_row(
        staff,
        row.pk,
        "UPDATE_SELECTED_EXISTING",
        "Revisi sumber dikonfirmasi",
        target_id=original.pk,
    )
    original.refresh_from_db()
    assert original.inflow_amount == 2000000 and original.revision_number == 2
    assert CashflowTransaction.objects.count() == 1


@pytest.mark.django_db
def test_replace_confirmation_scope_diff_and_soft_archive(superuser, staff, grant):
    first = stage_import(
        superuser,
        upload(workbook_bytes(source_rows({}, {"bank_name": "BANK B"}))),
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    commit_batch(superuser, first.pk)
    with pytest.raises(ValidationError):
        stage_import(superuser, upload(workbook_bytes(marker="replace")), mode="REPLACE")
    batch = stage_import(
        superuser,
        upload(workbook_bytes(marker="replace")),
        mode="REPLACE",
        scope=SCOPE,
        full_snapshot_confirmed=True,
        is_demo=True,
    )
    plan = preview_replace(superuser, batch.pk)
    assert (
        len(plan["deactivate"]) == 1
        and plan["before"]["inflow"] == "2000000"
        and plan["after"]["inflow"] == "1000000"
    )
    assert CashflowTransaction.objects.filter(soft_deleted_at__isnull=True).count() == 2
    with pytest.raises(ValidationError):
        commit_replace(superuser, batch.pk, plan["confirmation_token"] + "tamper", "Tes tamper")
    grant(staff, "import_data")
    staff_plan = preview_replace(staff, batch.pk)
    with pytest.raises(PermissionDenied):
        commit_replace(staff, batch.pk, staff_plan["confirmation_token"], "Tanpa izin arsip massal")
    commit_replace(
        superuser, batch.pk, plan["confirmation_token"], "Snapshot lengkap telah diperiksa"
    )
    assert CashflowTransaction.objects.filter(soft_deleted_at__isnull=True).count() == 1
    assert CashflowTransaction.objects.count() == 2


@pytest.mark.django_db
def test_replace_stale_preview_and_outside_scope_rejected(superuser):
    first = stage_import(
        superuser, upload(workbook_bytes()), is_demo=True, scope=SCOPE, full_snapshot_confirmed=True
    )
    commit_batch(superuser, first.pk)
    replacement = stage_import(
        superuser,
        upload(workbook_bytes(marker="replace")),
        mode="REPLACE",
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    plan = preview_replace(superuser, replacement.pk)
    archive_transactions(superuser, [CashflowTransaction.objects.get().pk], "Perubahan bersamaan")
    with pytest.raises(ValidationError):
        commit_replace(superuser, replacement.pk, plan["confirmation_token"], "Stale")
    outside = stage_import(
        superuser,
        upload(workbook_bytes(source_rows(transaction_date="2022-01-01"))),
        mode="REPLACE",
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    assert outside.stats == {"INVALID": 1}
    with pytest.raises(ValidationError):
        preview_replace(superuser, outside.pk)


@pytest.mark.django_db
def test_concurrent_stages_reconcile_again_at_commit(superuser):
    first = stage_import(superuser, upload(workbook_bytes(marker="first")), is_demo=True)
    second = stage_import(superuser, upload(workbook_bytes(marker="second")), is_demo=True)
    assert first.stats == second.stats == {"NEW_VALID": 1}
    commit_batch(superuser, first.pk)
    commit_batch(superuser, second.pk)
    assert CashflowTransaction.objects.count() == 1
    assert second.rows.get().status == "NEEDS_REVIEW"


@pytest.mark.django_db
def test_recoverable_batch_failure_rolls_back_all_rows(superuser, monkeypatch):
    from imports import services

    batch = stage_import(
        superuser, upload(workbook_bytes(source_rows({}, {"bank_name": "BANK B"}))), is_demo=True
    )
    original = services._create_or_update
    calls = []

    def fail_second(*args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            raise ValidationError("Simulasi gangguan")
        return original(*args, **kwargs)

    monkeypatch.setattr(services, "_create_or_update", fail_second)
    with pytest.raises(ValidationError):
        commit_batch(superuser, batch.pk)
    assert not CashflowTransaction.objects.exists()
    assert batch.rows.filter(status="NEW_VALID").count() == 2
    monkeypatch.setattr(services, "_create_or_update", original)
    commit_batch(superuser, batch.pk)
    assert CashflowTransaction.objects.count() == 2


@pytest.mark.django_db
def test_staging_failure_removes_private_source_file(superuser, monkeypatch, settings):
    def fail_insert(*args, **kwargs):
        raise ValidationError("Simulasi kegagalan staging database")

    monkeypatch.setattr(ImportRow.objects, "create", fail_insert)
    with pytest.raises(ValidationError):
        stage_import(superuser, upload(workbook_bytes()), is_demo=True)
    assert not ImportBatch.objects.exists()
    assert not list(Path(settings.MEDIA_ROOT).rglob("*.xlsx"))


@pytest.mark.django_db
def test_replace_selected_update_waits_for_confirmation(superuser):
    original_scope = {**SCOPE, "end": "2021-10-31"}
    original = stage_import(
        superuser,
        upload(workbook_bytes(source_rows({}, {"transaction_date": "2021-10-10"}))),
        is_demo=True,
        scope=original_scope,
        full_snapshot_confirmed=True,
    )
    commit_batch(superuser, original.pk)
    target = CashflowTransaction.objects.get(transaction_date="2021-09-10")
    incoming = stage_import(
        superuser,
        upload(
            workbook_bytes(
                source_rows(
                    source_amount=2000000, pieces=20, inflow_amount=2000000, source_netflow=2000000
                )
            )
        ),
        mode="REPLACE",
        is_demo=True,
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    row = incoming.rows.get()
    decide_row(
        superuser, row.pk, "UPDATE_SELECTED_EXISTING", "Sumber direvisi", target_id=target.pk
    )
    target.refresh_from_db()
    assert target.inflow_amount == 1000000
    row.refresh_from_db()
    assert row.status == "READY_UPDATE"
    plan = preview_replace(superuser, incoming.pk)
    assert len(plan["update"]) == 1 and plan["after"]["inflow"] == "2000000"
    commit_replace(superuser, incoming.pk, plan["confirmation_token"], "Konfirmasi revisi snapshot")
    target.refresh_from_db()
    assert target.inflow_amount == 2000000
    assert (
        CashflowTransaction.objects.filter(
            transaction_date="2021-10-10", soft_deleted_at__isnull=True
        ).count()
        == 1
    )


@pytest.mark.django_db(transaction=True)
def test_two_parallel_commits_do_not_duplicate(superuser):
    batches = [
        stage_import(superuser, upload(workbook_bytes(marker=str(index))), is_demo=True)
        for index in range(2)
    ]

    def worker(batch_id):
        close_old_connections()
        try:
            user = get_user_model().objects.get(pk=superuser.pk)
            return commit_batch(user, batch_id).status
        finally:
            close_old_connections()

    with ThreadPoolExecutor(max_workers=2) as pool:
        statuses = list(pool.map(worker, [batch.pk for batch in batches]))
    assert sorted(statuses) == ["COMMITTED", "PARTIAL"]
    assert CashflowTransaction.objects.count() == 1
    assert ImportRow.objects.filter(status="NEEDS_REVIEW").count() == 1


@pytest.mark.parametrize("label", ["1.000 UK", "1.000 UL", "  1.000 uk  "])
def test_setbay_denomination_labels_in_reference_and_transaction(label):
    from openpyxl import load_workbook

    data = workbook_bytes(
        source_rows(
            denomination_value=label,
            source_amount=10000,
            inflow_amount=10000,
            source_netflow=10000,
        )
    )
    book = load_workbook(io.BytesIO(data))
    book["DENOM"]["B2"] = label
    output = io.BytesIO()
    book.save(output)
    parsed = parse_workbook(output.getvalue(), "synthetic.xlsx")
    row = parsed.rows[0]
    assert row.status == "NEW_VALID"
    assert row.normalized["denomination_value"] == "1000"
    assert row.raw["denomination_value"]["value"] == label
    assert parsed.denomination_references["SYN100"]["value"] == "1000"


@pytest.mark.parametrize(
    "code,reference,expected",
    [
        (760002, 760002, "NEW_VALID"),
        (1, "000001", "INVALID"),
    ],
)
def test_numeric_kastip_matches_reference_without_guessing_leading_zeros(code, reference, expected):
    book = Workbook()
    sheet = book.active
    sheet.title = "SETBAY KASTIP"
    keys = list(ADAPTERS[sheet.title])
    sheet.append([ADAPTERS[sheet.title][key][0] for key in keys])
    row = source_rows(region_code=code, region_label="SKM SINTETIS")[0]
    sheet.append([row[key] for key in keys])
    refs = book.create_sheet("DENOM")
    refs.append(["Pecahan", "Denom", "Urutan Denom", None, "Kode Kastip", "Nama Kastip"])
    refs.append(["SYN100", "100.000", 1, None, reference, "SKM SINTETIS"])
    output = io.BytesIO()
    book.save(output)
    parsed = parse_workbook(output.getvalue(), "synthetic.xlsx")
    assert parsed.rows[0].status == expected
    assert parsed.rows[0].normalized["region_code"] == str(code)


def test_numeric_account_is_canonical_with_nonblocking_warning():
    row = parse_workbook(workbook_bytes(source_rows(account=12345.0)), "synthetic.xlsx").rows[0]
    assert row.normalized["account"] == "12345"
    assert row.status == "NEW_VALID"
    assert "numeric_account" in {issue["code"] for issue in row.normalized["warnings"]}


def test_denomination_suffix_is_not_accepted_on_amounts():
    row = parse_workbook(
        workbook_bytes(source_rows(source_amount="1.000 UK")), "synthetic.xlsx"
    ).rows[0]
    assert row.status == "INVALID"
