import io
from decimal import Decimal

import pytest
from django.core.exceptions import PermissionDenied
from openpyxl import load_workbook

from cashflow.models import CashflowTransaction
from core.models import Denomination, ServiceRegion, SystemSettings
from imports.parser import ADAPTERS, WorkbookError, parse_workbook
from imports.services import import_summary, smart_import
from tests.test_imports import SCOPE, source_rows, upload, workbook_bytes


def two_sheet_workbook():
    book = load_workbook(io.BytesIO(workbook_bytes()))
    sheet = book.create_sheet("SETBAY KASTIP")
    keys = list(ADAPTERS[sheet.title])
    sheet.append([ADAPTERS[sheet.title][key][0] for key in keys])
    row = source_rows(region_code=760002, region_label="SKM SINTETIS", account=12345)[0]
    sheet.append([row[key] for key in keys])
    book["DENOM"]["E2"] = 760002
    # Reorder references: mapping must follow column labels.
    refs = book["DENOM"]
    rows = list(refs.values)
    book.remove(refs)
    refs = book.create_sheet("DENOM")
    for row in rows:
        refs.append(list(reversed(row)))
    book.create_sheet("Tanggal Update").append(["Tanggal Update"])
    book["Tanggal Update"].append(["2026-03-26"])
    book.create_sheet("Sheet2")
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


@pytest.mark.django_db
def test_both_sheets_smart_import_numeric_warning_and_same_file(superuser):
    content = two_sheet_workbook()
    batch = smart_import(superuser, upload(content))
    assert import_summary(batch) == dict(
        total=2, saved=2, existing=0, warnings=1, review=0, failed=0
    )
    assert batch.metadata["source_update_date"] == "2026-03-26"
    assert CashflowTransaction.objects.get(region__code="760002").account == "12345"
    assert CashflowTransaction.objects.get(region__code="MANADO").account == "0000000001"
    assert sum(CashflowTransaction.objects.values_list("inflow_amount", flat=True)) == 2000000
    assert SystemSettings.load().current_publication_version_id is None
    again = smart_import(superuser, upload(content, "renamed.xlsx"))
    assert again.pk == batch.pk
    assert CashflowTransaction.objects.count() == 2


@pytest.mark.django_db
def test_distinct_valid_occurrences_in_same_upload_are_saved(superuser):
    batch = smart_import(
        superuser,
        upload(
            workbook_bytes(
                source_rows(
                    {},
                    {
                        "source_amount": 2000000,
                        "pieces": 20,
                        "inflow_amount": 2000000,
                        "source_netflow": 2000000,
                    },
                )
            )
        ),
    )
    assert batch.stats == {"IMPORTED": 2}
    assert CashflowTransaction.objects.count() == 2


@pytest.mark.django_db
def test_partial_with_warning_ambiguous_and_invalid(superuser):
    batch = smart_import(
        superuser,
        upload(
            workbook_bytes(
                source_rows(
                    {"account": 12345},
                    {"account": 12345},
                    {"bank_name": "BANK B", "pieces": 9},
                    {"bank_name": "BANK C", "transaction_date": "bad"},
                )
            )
        ),
    )
    assert batch.stats == {"IMPORTED": 1, "NEEDS_REVIEW": 2, "INVALID": 1}
    assert import_summary(batch)["warnings"] == 2


@pytest.mark.django_db
def test_existing_master_conflict_and_region_conflict_hold_only_affected_row(superuser):
    Denomination.objects.create(code="SYN100", value=50000, verified=True)
    batch = smart_import(superuser, upload(workbook_bytes()))
    assert batch.stats == {"NEEDS_REVIEW": 1}
    assert Denomination.objects.get(code="SYN100").value == Decimal(50000)
    Denomination.objects.filter(code="SYN100").update(value=100000)
    ServiceRegion.objects.create(code="MANADO", name="Nama berbeda", kind="MANADO")
    batch = smart_import(superuser, upload(workbook_bytes(marker="region conflict")))
    assert batch.stats == {"NEEDS_REVIEW": 1}
    assert "region_master_conflict" in {x["code"] for x in batch.rows.get().validation_errors}


@pytest.mark.django_db
@pytest.mark.parametrize("first_account,next_account", [(12345, "0012345"), ("0012345", 12345)])
def test_uncertain_account_never_silently_duplicates_history(
    superuser, first_account, next_account
):
    smart_import(superuser, upload(workbook_bytes(source_rows(account=first_account))))
    incoming = smart_import(superuser, upload(workbook_bytes(source_rows(account=next_account))))
    assert incoming.stats == {"NEEDS_REVIEW": 1}
    assert CashflowTransaction.objects.count() == 1


@pytest.mark.django_db
def test_confirmed_history_skipped_and_new_scope_rows_saved(superuser):
    smart_import(superuser, upload(workbook_bytes()), scope=SCOPE, full_snapshot_confirmed=True)
    batch = smart_import(
        superuser,
        upload(workbook_bytes(source_rows({}, {"bank_name": "BANK BARU"}), marker="next")),
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    assert batch.stats == {"ALREADY_IMPORTED": 1, "IMPORTED": 1}


@pytest.mark.django_db
def test_update_replace_and_permissions_still_require_explicit_action(superuser, staff):
    with pytest.raises(PermissionDenied):
        smart_import(staff, upload(workbook_bytes()))
    update = smart_import(superuser, upload(workbook_bytes()), mode="UPDATE")
    assert update.stats == {"NEEDS_REVIEW": 1}
    replace = smart_import(
        superuser,
        upload(workbook_bytes(marker="replace")),
        mode="REPLACE",
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    assert replace.stats == {"NEW_VALID": 1}
    assert not CashflowTransaction.objects.exists()


@pytest.mark.parametrize("account", [1234567890123456, -123, 1.5])
def test_numeric_account_precision_risk_remains_blocking(account):
    row = parse_workbook(workbook_bytes(source_rows(account=account)), "synthetic.xlsx").rows[0]
    assert row.status in {"NEEDS_REVIEW", "INVALID"}


def test_conflicting_reference_is_rejected():
    book = load_workbook(io.BytesIO(workbook_bytes()))
    book["DENOM"].append(["SYN100", "50.000", 2])
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(WorkbookError):
        parse_workbook(out.getvalue(), "synthetic.xlsx")


@pytest.mark.django_db
def test_upload_ui_saves_without_second_commit_and_hides_raw_dict(client, superuser):
    client.force_login(superuser)
    response = client.post(
        "/admin/imports/", {"file": upload(two_sheet_workbook()), "mode": "ADD"}, follow=True
    )
    assert response.status_code == 200
    assert CashflowTransaction.objects.count() == 2
    assert b"Berhasil disimpan" in response.content
    assert b"numeric_account" not in response.content
    assert b"NEW_VALID" not in response.content
    assert b"&#x27;code&#x27;" not in response.content


def test_functional_header_aliases_do_not_double_count():
    book = load_workbook(io.BytesIO(workbook_bytes()))
    sheet = book["SETBAY MANADO"]
    for cell in sheet[1]:
        if isinstance(cell.value, str):
            cell.value = "  " + cell.value.replace("COTCP_", "COTCR_").lower() + "  "
    out = io.BytesIO()
    book.save(out)
    row = parse_workbook(out.getvalue(), "synthetic.xlsx").rows[0]
    assert row.status == "NEW_VALID"
    assert row.normalized["inflow_amount"] == "1000000"
    assert row.normalized["outflow_amount"] == "0"


@pytest.mark.django_db
def test_replace_ui_keeps_confirmation_and_shows_human_summary(client, superuser):
    batch = smart_import(
        superuser,
        upload(workbook_bytes()),
        mode="REPLACE",
        scope=SCOPE,
        full_snapshot_confirmed=True,
    )
    client.force_login(superuser)
    response = client.post(f"/admin/imports/{batch.pk}/", {"action": "replace_preview"})
    assert response.status_code == 200
    assert b"Sebelum" in response.content and b"Sesudah" in response.content
    assert b"confirmation_token" in response.content
    assert b"&quot;batch_id&quot;" not in response.content
    assert not CashflowTransaction.objects.exists()


def test_irrelevant_sheet_is_ignored_and_sheet_name_is_normalized():
    book = load_workbook(io.BytesIO(workbook_bytes()))
    book["SETBAY MANADO"].title = " setbay   manado "
    unrelated = book.create_sheet("Catatan lain")
    unrelated.cell(row=100000, column=200, value="Tidak terkait transaksi")
    out = io.BytesIO()
    book.save(out)
    parsed = parse_workbook(out.getvalue(), "synthetic.xlsx")
    assert len(parsed.rows) == 1
    assert parsed.rows[0].status == "NEW_VALID"
